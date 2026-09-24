#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
filter_maoli_ljc.py —— 把经理号导出的「全公司」毛利明细过滤成「李家村」专用明细。
（storyos B 方案组件，2026-09-24 建；镜像 shophyc/filter_maoli_hyc.py，仅换门店键）

背景：storyos 双店统一用经理工号(18591910491)拉数，导出的是全公司 17 店明细。
不过滤直接喂 calc_data.py / write_xs_xml.py 会把 17 家店当成一家复算 → 数字全部虚高。

过滤口径（同华阳城版）：保留「库区」或「销售出库单门店」包含 "李家村" 的行。

用法：
    python filter_maoli_ljc.py <源明细.xlsx> [输出.xlsx]
"""
import os
import sys
import warnings

from openpyxl import load_workbook, Workbook

STORE_KEY = "李家村"
HEADERS = ["出库单号", "单据类型", "出库日期", "商品分类", "商品sku分类", "商品SKU编码",
           "商品名称", "入库属性", "数量", "单价", "原价", "折扣价", "金额", "毛利",
           "SO激励", "业务员", "库区", "销售出库单门店", "销售成本"]
DEF_OUT_DIR = "/Users/mac/.local/share/TeleAgent/playwright-mcp/ljc"
DEF_OUT = os.path.join(DEF_OUT_DIR, "李家村门店毛利明细表-华为终端.xlsx")


def main():
    if len(sys.argv) < 2:
        sys.exit("用法: filter_maoli_ljc.py <源明细.xlsx> [输出.xlsx]")
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else DEF_OUT
    if not os.path.exists(src):
        sys.exit(f"❌ 找不到源明细: {src}")

    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    if hdr[:19] != HEADERS:
        sys.exit(f"❌ 源明细表头不符（期望 19 列）\n  实际: {hdr[:19]}")
    i_kq, i_md = 16, 17   # 库区 / 销售出库单门店

    keep, total = [], 0
    for r in it:
        if not r or not r[0] or not str(r[0]).strip():
            continue
        total += 1
        kq = STORE_KEY in str(r[i_kq] or "")
        md = STORE_KEY in str(r[i_md] or "")
        if kq or md:
            keep.append(r[:19])

    if not keep:
        sys.exit(f"❌ 过滤后 0 行（源 {total} 行中无「{STORE_KEY}」）—— 检查门店名")

    os.makedirs(os.path.dirname(out), exist_ok=True)
    nwb = Workbook(write_only=True)
    nws = nwb.create_sheet("门店毛利明细表-华为终端")
    nws.append(HEADERS)
    for r in keep:
        nws.append(["" if v is None else v for v in r])
    nwb.save(out)

    # 校验：行数 + 毛利合计 + 人员分布（混入他店员工 = 过滤失败信号）
    import collections
    emp = collections.Counter(str(r[15] or "") for r in keep)
    gross = 0.0
    for r in keep:
        try:
            gross += float(str(r[13]).replace(",", ""))
        except (TypeError, ValueError):
            pass
    print(f"→ 源 {total} 行 → 保留「{STORE_KEY}」{len(keep)} 行")
    print(f"  毛利合计 ¥{gross:,.2f}")
    print(f"  业务员分布: {dict(emp)}")
    print(f"✓ 已写出: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
